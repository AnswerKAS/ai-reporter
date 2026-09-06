"""Отчёт файлом: Excel и PDF (reports/render.py)."""

import zipfile
from io import BytesIO

import pytest

from app.reports import render

REPORT = {
    'slug': 'sales', 'title': 'Продажи', 'description': 'За январь',
    'filters': [{'key': 'city', 'label': 'Город'}],
    'filterValues': {'city': 'Москва'},
    'sections': [
        {'type': 'kpi', 'items': [{'label': 'Выручка', 'value': 1234.5, 'format': 'money'}]},
        {'type': 'chart', 'kind': 'bar', 'title': 'По дням', 'xKey': 'day',
         'series': [{'key': 'revenue', 'name': 'Выручка'}],
         'data': [{'day': '2026-01-01', 'revenue': 10}, {'day': '2026-01-02', 'revenue': 20}]},
        {'type': 'table', 'title': 'По городам',
         'columns': [{'key': 'city', 'header': 'Город'},
                     {'key': 'revenue', 'header': 'Выручка', 'format': 'money'}],
         'rows': [{'city': 'Москва', 'revenue': 100}]},
    ],
}


# --- форматирование значений ------------------------------------------------

@pytest.mark.parametrize('value,fmt,expected', [
    (None, None, ''),
    (1234.5, None, '1 234.5'),
    (1234.5, 'money', '1 234.5 ₽'),
    (12.0, 'percent', '12 %'),
    (1000, None, '1 000'),
    ('текст', None, 'текст'),
    (True, None, 'True'),
])
def test_форматирование(value, fmt, expected):
    assert render._fmt(value, fmt) == expected


# --- разбор секций в таблицу --------------------------------------------------

def test_карточки_разворачиваются_в_две_колонки():
    head, rows = render._section_rows(REPORT['sections'][0])

    assert head == ['Показатель', 'Значение']
    assert rows == [['Выручка', 1234.5]]


def test_график_разворачивается_по_сериям():
    head, rows = render._section_rows(REPORT['sections'][1])

    assert head == ['day', 'Выручка']
    assert rows == [['2026-01-01', 10], ['2026-01-02', 20]]


def test_таблица_берёт_колонки_из_описания():
    head, rows = render._section_rows(REPORT['sections'][2])

    assert head == ['Город', 'Выручка']
    assert rows == [['Москва', 100]]


def test_неизвестная_секция_даёт_пусто():
    assert render._section_rows({'type': 'markdown', 'content': 'x'}) == ([], [])


def test_заголовок_секции_по_умолчанию():
    assert render._title({'type': 'kpi'}, 1) == 'Показатели'
    assert render._title({'type': 'chart'}, 1) == 'График'
    assert render._title({'type': 'table', 'title': 'Своё'}, 1) == 'Своё'
    assert render._title({'type': 'markdown'}, 3) == 'Секция 3'


# --- Excel ------------------------------------------------------------------------

def sheets(data: bytes) -> list[str]:
    from openpyxl import load_workbook

    return load_workbook(BytesIO(data)).sheetnames


def test_excel_лист_на_секцию():
    names = sheets(render.to_xlsx(REPORT))

    assert names[0] == 'Отчёт'
    assert 'По дням' in names and 'По городам' in names


def test_excel_первый_лист_это_сводка():
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(render.to_xlsx(REPORT)))
    values = [row[0] for row in book['Отчёт'].values]

    assert values[0] == 'Продажи'
    assert 'Город' in values  # применённый фильтр


def test_excel_числа_остаются_числами():
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(render.to_xlsx(REPORT)))

    assert book['По городам']['B2'].value == 100


def test_excel_имена_листов_уникальны_и_коротки():
    report = {'slug': 'r', 'sections': [
        {'type': 'table', 'title': 'Очень длинное название секции отчёта',
         'columns': [{'key': 'a', 'header': 'A'}], 'rows': []},
        {'type': 'table', 'title': 'Очень длинное название секции отчёта',
         'columns': [{'key': 'a', 'header': 'A'}], 'rows': []},
    ]}

    names = sheets(render.to_xlsx(report))[1:]

    assert len(set(names)) == 2
    assert all(len(n) <= 31 for n in names)


def test_excel_отчёта_без_секций():
    assert sheets(render.to_xlsx({'slug': 'r', 'sections': []})) == ['Отчёт']


# --- PDF ---------------------------------------------------------------------------

def test_pdf_собирается():
    data = render.to_pdf(REPORT)

    assert data[:5] == b'%PDF-'
    assert len(data) > 2000


def test_pdf_пустого_отчёта():
    data = render.to_pdf({'slug': 'r', 'title': 'Пустой', 'sections': []})

    assert data[:5] == b'%PDF-'


def test_pdf_обрезает_длинную_таблицу():
    report = {'slug': 'r', 'title': 'Большой', 'sections': [
        {'type': 'table', 'title': 'Много строк',
         'columns': [{'key': 'a', 'header': 'A'}],
         'rows': [{'a': i} for i in range(render.PDF_TABLE_ROWS + 10)]}]}

    assert render.to_pdf(report)[:5] == b'%PDF-'


@pytest.mark.parametrize('kind', ['bar', 'line', 'area', 'pie'])
def test_картинка_графика_рисуется(kind):
    section = {**REPORT['sections'][1], 'kind': kind}

    assert render._chart_image(section)[:4] == b'\x89PNG'


def test_график_без_данных_картинки_не_даёт():
    assert render._chart_image({'type': 'chart', 'data': [], 'series': []}) is None


# --- выбор формата -------------------------------------------------------------------

def test_рендер_xlsx_по_умолчанию():
    data, name, mime = render.render(REPORT, 'xlsx')

    assert name == 'sales.xlsx'
    assert mime.endswith('spreadsheetml.sheet')
    assert zipfile.is_zipfile(BytesIO(data))


def test_рендер_pdf():
    data, name, mime = render.render(REPORT, 'pdf')

    assert (name, mime) == ('sales.pdf', 'application/pdf')


def test_неизвестный_формат_даёт_excel():
    _, name, _ = render.render(REPORT, 'docx')

    assert name.endswith('.xlsx')
